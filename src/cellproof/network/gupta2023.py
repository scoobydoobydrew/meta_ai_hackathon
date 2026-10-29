"""Build the base network of Gupta et al. 2023 from the rule rows of Supplementary Table S2.

Gupta S, Panda PK, Silveira DA, Ahuja R, Hashimoto RF. Cells 2023;12(7):1085, doi 10.3390/cells12071085 (CC BY 4.0).

Table S2 has one "Rule" row per node. The paper spells nodes several ways (ATM/pH2AX, ATM-pH2AX,
ATM_pH2AX; mTOR2 for mTORC2; ULK1 for ULK1/Beclin1), so every identifier is normalised and must
resolve to one of NODES. Anything else raises instead of being guessed.

Usage: python -m cellproof.network.gupta2023 <cells-12-01085-s001.zip> <out.bnet>
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

INPUT = "DNA_damage"
NODES = [
    INPUT, "ATM_pH2AX", "AMPK_MAPK", "Mdm2", "p53", "p53_A", "p53_INP1", "p53_K", "p21", "Wip1",
    "RB", "Myc", "WWP1", "hTERT", "E2F1", "Cdc25", "Sirt1", "Cdc2_CycB", "Wee1", "TUG1", "miR221",
    "PTEN", "AKT", "mTORC1", "mTORC2", "ULK1_Beclin1", "BCL2", "PUMA", "BAX", "Caspase",
]

# Table S2 drops the last ")" of the mTORC2 rule; Table S1 prints the same rule balanced.
S2_MISSING_CLOSING_PAREN = {"mTORC2"}

_OPS = {"AND": "&", "OR": "|", "NOT": "!"}


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


ALIASES = {_key(n): n for n in NODES} | {"mtor1": "mTORC1", "mtor2": "mTORC2", "ulk1": "ULK1_Beclin1"}


def translate(rule: str) -> str:
    """Rewrite `A AND NOT (B OR C)` as `A & !(B | C)`, resolving every name to a node id."""
    out, depth = "", 0
    for tok in re.findall(r"\(|\)|[^\s()]+", rule):
        if tok in _OPS:
            out += f" {_OPS[tok]} " if tok != "NOT" else "!"
        elif tok in "()":
            depth += 1 if tok == "(" else -1
            if depth < 0:
                raise ValueError(f"unbalanced parentheses in {rule!r}")
            out += tok
        elif _key(tok) in ALIASES:
            out += ALIASES[_key(tok)]
        else:
            raise ValueError(f"unknown identifier {tok!r} in rule {rule!r}")
    if depth != 0:
        raise ValueError(f"unbalanced parentheses in {rule!r}")
    return out


def read_rules(supplement_zip: str | Path) -> dict[str, str]:
    """{node id: raw rule text} from the "Rule" rows of Table S2."""
    with zipfile.ZipFile(supplement_zip) as z:
        wb = openpyxl.load_workbook(io.BytesIO(z.read("Table S2.xlsx")), read_only=True)
    rules: dict[str, str] = {}
    target = None
    header_seen = False
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if row[4] == "Target Node":
            header_seen = str(row[5]).startswith("No. Interactions")
            continue
        if not header_seen:
            continue
        if row[4]:
            target = ALIASES[_key(str(row[4]))]
        if row[5] == "Rule":
            if target in rules:
                raise ValueError(f"two rules for {target}")
            rules[target] = str(row[6]).strip()
    if not header_seen:
        raise ValueError("Table S2 header row not found; the layout changed")
    return rules


def read_annotations(supplement_zip: str | Path) -> dict:
    """The official molecule names and the biological justification of every interaction, as Table S2 prints them:
    {"nodes": {node id: official name}, "interactions": [{target, regulator, text, references}]}."""
    with zipfile.ZipFile(supplement_zip) as z:
        wb = openpyxl.load_workbook(io.BytesIO(z.read("Table S2.xlsx")), read_only=True)
    names: dict[str, str] = {}
    interactions: list[dict] = []
    target, header_seen = None, False
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if row[4] == "Target Node":
            header_seen = True
            continue
        if not header_seen:
            continue
        if row[4]:
            target = ALIASES[_key(str(row[4]))]
            if row[1]:
                names[target] = " ".join(str(row[1]).split())
        if isinstance(row[5], int | float) and row[6] and row[7]:
            interactions.append({
                "target": target, "regulator": ALIASES[_key(str(row[6]))],
                "text": " ".join(str(row[7]).split()), "references": " ".join(str(row[8] or "").split()),
            })
    return {"source": "Gupta et al. 2023, Cells 12(7):1085, Supplementary Table S2 (article CC BY 4.0)", "nodes": names, "interactions": interactions}


def build(raw_rules: dict[str, str]) -> dict[str, str]:
    """Validated {node id: bnet rule}, including the constant input node."""
    rules = {INPUT: INPUT}
    for node, raw in raw_rules.items():
        if node in S2_MISSING_CLOSING_PAREN and raw.count("(") - raw.count(")") == 1:
            raw += ")"
        rules[node] = translate(raw)
    if set(rules) != set(NODES):
        raise ValueError(f"missing {set(NODES) - set(rules)}, unexpected {set(rules) - set(NODES)}")
    return {n: rules[n] for n in NODES}


def to_bnet(rules: dict[str, str]) -> str:
    header = (
        "# Gupta et al. 2023, Cells 12(7):1085, doi 10.3390/cells12071085 (article CC BY 4.0).\n"
        "# Rules read from the Rule rows of Supplementary Table S2 by cellproof.network.gupta2023;\n"
        "# the mTORC2 rule lacked its last closing parenthesis in Table S2 and is balanced as in Table S1.\n"
        "# DNA_damage is the model's only input (constant). Node ids differ from the paper's spelling for:\n"
        "# ATM_pH2AX (ATM/pH2AX), AMPK_MAPK (AMPK/MAPK), p53_A, p53_K, p53_INP1, Cdc2_CycB, miR221 (miR-221),\n"
        "# ULK1_Beclin1 (ULK1/Beclin1; the S2 regulator list says ULK1), mTORC1/mTORC2 (S2 also says mTOR1/mTOR2).\n"
        "targets, factors\n"
    )
    return header + "".join(f"{n}, {r}\n" for n, r in rules.items())


if __name__ == "__main__":
    zip_path, out_path = sys.argv[1:3]
    Path(out_path).write_text(to_bnet(build(read_rules(zip_path))), encoding="utf-8", newline="\n")
    print("wrote", out_path)
